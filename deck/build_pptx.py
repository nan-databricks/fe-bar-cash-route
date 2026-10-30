#!/usr/bin/env python3
"""Builds deck/SNB_Cash_Logistics.pptx from the designed HTML deck.

Rebuilds deck/deck.pdf, renders each page to PNG, places it full-bleed on a
slide, sets a real (searchable) slide title hidden behind the image, and writes
speaker notes from deck/speaker_notes.json plus docs/demo_spoken_script.md.

Usage: python3 deck/build_pptx.py   (needs python-pptx and poppler's pdftoppm)
"""
import html
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches

DECK = Path(__file__).resolve().parent
ROOT = DECK.parent
PDF = DECK / "deck.pdf"
OUT = DECK / "SNB_Cash_Logistics.pptx"
NOTES_JSON = DECK / "speaker_notes.json"
SCRIPT_MD = ROOT / "docs" / "demo_spoken_script.md"
N_SLIDES = 14


def build_pdf():
    subprocess.run(["bash", str(DECK / "build_pdf.sh")], check=True)
    info = subprocess.run(["pdfinfo", str(PDF)], capture_output=True, text=True, check=True).stdout
    pages = int(re.search(r"Pages:\s+(\d+)", info).group(1))
    if pages != N_SLIDES:
        sys.exit(f"deck.pdf has {pages} pages, expected {N_SLIDES}")


def render_pngs(tmp):
    subprocess.run(["pdftoppm", "-png", "-r", "200", str(PDF), str(tmp / "p")], check=True)
    pngs = sorted(tmp.glob("p-*.png"), key=lambda p: int(p.stem.split("-")[1]))
    if len(pngs) != N_SLIDES:
        sys.exit(f"rendered {len(pngs)} PNGs, expected {N_SLIDES}")
    return pngs


def html_titles():
    src = (DECK / "slides.html").read_text()
    sections = re.findall(r'<section class="slide[^"]*".*?</section>', src, re.S)
    titles = []
    for sec in sections:
        m = re.search(r"<h[12][^>]*>(.*?)</h[12]>", sec, re.S)
        text = ""
        if m:
            text = re.sub(r"<br\s*/?>", " ", m.group(1))
            text = html.unescape(re.sub(r"<[^>]+>", "", text))
            text = re.sub(r"\s+", " ", text).strip()
        titles.append(text)
    return titles


# ---------- spoken script parsing ----------

def plain(text):
    text = text.replace(r"\*\*", "asterisk")
    text = re.sub(r"\*\*|\*|`", "", text)
    return re.sub(r"\s+", " ", text).strip()


def classify(lines):
    """Split markdown lines into (say paragraphs, cue lines)."""
    say, cue = [], []
    for raw in lines:
        line = raw.strip()
        if not line or line == "---" or line.startswith("#"):
            continue
        if line.startswith("**Click:**"):
            cue.append("Click: " + plain(line[len("**Click:**"):]))
            continue
        if line.startswith("Say:"):
            say.append(plain(line[4:]))
            continue
        if line.startswith("*"):
            # Italic segments are stage directions; any non-italic remainder is spoken.
            italics = re.findall(r"(?<!\*)\*([^*]+)\*(?!\*)", line)
            rest = re.sub(r"(?<!\*)\*[^*]+\*(?!\*)", "", line)
            for seg in italics:
                seg = plain(seg).rstrip(":")
                if not re.fullmatch(r"Slide \d+\.?", seg):
                    cue.append(seg)
            if plain(rest):
                say.append(plain(rest))
            continue
        say.append(plain(line))
    return say, cue


def script_sections():
    md = SCRIPT_MD.read_text()
    parts = re.split(r"^## (\d+)\. (.*)$", md, flags=re.M)
    secs = {}
    for i in range(1, len(parts), 3):
        num, heading, body = int(parts[i]), parts[i + 1], parts[i + 2]
        body = re.split(r"^## ", body, flags=re.M)[0]
        t = re.search(r"\((\d+:\d+)[–-](\d+:\d+)\)", heading)
        secs[num] = {
            "heading": plain(heading),
            "time": f"{t.group(1)}–{t.group(2)}" if t else "",
            "lines": body.splitlines(),
        }
    return secs


def split_at(lines, markers):
    """Split lines before each line that starts with one of the markers (in order)."""
    chunks, cur, mi = [], [], 0
    for line in lines:
        if mi < len(markers) and line.strip().startswith(markers[mi]):
            chunks.append(cur)
            cur, mi = [], mi + 1
        cur.append(line)
    chunks.append(cur)
    if len(chunks) != len(markers) + 1:
        sys.exit(f"script markers not found: {markers}")
    return chunks


def script_for_slides(secs):
    """Map slide number -> dict(intro, say, cue, time)."""
    out = {}
    s1 = split_at(secs[1]["lines"], ["*Slide 2.*", "*Slide 3.*", "*Slide 4.*"])
    for n, chunk in enumerate(s1, start=1):
        say, cue = classify(chunk)
        out[n] = {"intro": None, "say": say, "cue": cue, "time": f"{secs[1]['time']} (section 1, setup)"}

    say, cue = [], []
    for k in (2, 3, 4):
        s, c = classify(secs[k]["lines"])
        say += [f"[Section {k}: {secs[k]['heading']}]"] + s
        cue += c
    out[5] = {"intro": "Backup if the app fails — say the section 2–4 lines while showing this.",
              "say": say, "cue": cue,
              "time": f"{secs[2]['time'].split('–')[0]}–{secs[4]['time'].split('–')[1]} (sections 2–4, app demo)"}

    s, c = classify(secs[5]["lines"])
    out[6] = {"intro": "Backup if the app fails — say the section 5 lines while showing this.",
              "say": s, "cue": c, "time": f"{secs[5]['time']} (section 5, Ask SNB Ops)"}

    s6 = split_at(secs[6]["lines"], ["**Click:** Lakeflow", "**Click:** MLflow"])
    for n, chunk in zip((7, 8, 9), s6):
        s, c = classify(chunk)
        out[n] = {"intro": "Backup for the workspace tabs — say the section 6 lines while showing this.",
                  "say": s, "cue": c, "time": f"{secs[6]['time']} (section 6, under the hood)"}

    s7 = split_at(secs[7]["lines"], ["*Slide 11.*"])
    for n, chunk in zip((10, 11), s7):
        s, c = classify(chunk)
        out[n] = {"intro": None, "say": s, "cue": c, "time": f"{secs[7]['time']} (section 7, close)"}

    for n in (12, 13, 14):
        out[n] = {"intro": None, "say": [], "cue": ["Appendix: use in Q&A"],
                  "time": "12:00–25:00 (Q&A window, as needed)"}
    return out


def notes_text(json_notes, s):
    lines = ["SAY:", plain(json_notes)]
    if s["intro"] or s["say"]:
        lines += ["", "From the spoken script:"]
        if s["intro"]:
            lines.append(s["intro"])
        lines += s["say"]
    lines += ["", "CLICK / CUE:"] + (s["cue"] or ["None, stay on the slide."])
    lines += ["", "TIME: " + s["time"]]
    return "\n".join(lines)


def main():
    build_pdf()
    notes = json.loads(NOTES_JSON.read_text())
    titles = html_titles()
    if len(notes) != N_SLIDES or len(titles) != N_SLIDES:
        sys.exit(f"expected {N_SLIDES} notes/titles, got {len(notes)}/{len(titles)}")
    per_slide = script_for_slides(script_sections())

    tmp = Path(tempfile.mkdtemp(prefix="pptx_src_"))
    try:
        pngs = render_pngs(tmp)
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        layout = next(l for l in prs.slide_layouts if l.name == "Title Only")
        for i, png in enumerate(pngs):
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = titles[i] or notes[i]["title"]
            # Added after the title, so the picture sits on top and hides it.
            slide.shapes.add_picture(str(png), Emu(0), Emu(0), prs.slide_width, prs.slide_height)
            slide.notes_slide.notes_text_frame.text = notes_text(notes[i]["notes"], per_slide[i + 1])
        prs.save(OUT)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print(f"wrote {OUT} ({OUT.stat().st_size / 1e6:.1f} MB)")


if __name__ == "__main__":
    main()
